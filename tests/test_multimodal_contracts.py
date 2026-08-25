"""多模态知识库的无网络契约与隔离行为测试。"""

from __future__ import annotations

import tempfile
import unittest
import os
import json
import io
import base64
import sqlite3
import threading
import time
from unittest.mock import MagicMock, patch
from pathlib import Path
from unittest.mock import patch

from PIL import Image

from Agent.knowledge_base.multimodal.assets import AssetStore
from Agent.knowledge_base.multimodal.index import StagedIndex
from Agent.knowledge_base.multimodal.contracts import BoundingBox, IngestionIssue, IssueSeverity, KnowledgeUnit, OutboundImageRecord, UnitStatus, VisionAnalysis, render_retrieval_text, sha256_bytes, stable_id
from Agent.knowledge_base.multimodal.pipeline import MultimodalKnowledgeBaseMaintenance
from Agent.knowledge_base.multimodal.parsers import ParsedDocument, ParsedItem, OcrResult, _convert_docling_page, _rapidocr_probe, decide_page_route, inspect_source, ocr_fingerprint, parse_document
from Agent.knowledge_base.rag_runtime import RagRuntimeConfig
from Agent.knowledge_base.multimodal.vision import PROMPT_VERSION, REQUIRED_MODEL, RESPONSE_ADAPTER_VERSION, VisionAnalyzer
from Agent.knowledge_base.query_rag import RagRetrievalConfig, _build_retrieval_trace_with_resources, _merge_candidates
from Agent.knowledge_base.sparse_retriever import candidate_identity
from Agent.tool_node.rag_questions import normalize_rag_question_output
from Agent.tool_node.rag_tool_registry import build_rag_tools


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]


def _vision_payload(**overrides: object) -> str:
    """生成完整的 vision-v2 fake 响应。"""
    payload = {"content_kind": "chart", "ocr_text": "图中文字", "visible_facts": ["存在坐标轴"], "summary": "图表摘要", "entities": ["A"], "table_markdown": "", "formula_latex": "", "directed_relations": [], "uncertain_relations": [], "confidence": 0.8, "informative": True}
    payload.update(overrides)
    return json.dumps(payload, ensure_ascii=False)


def _png_bytes(color: str = "white", size: tuple[int, int] = (8, 8)) -> bytes:
    """生成无需外部 fixture 的有效 PNG。"""
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


def _outbound_record(analyzer: VisionAnalyzer, payload: bytes, context: str = "", image_index: int = 1, *, source_relative_path: str = "source.pdf", source_sha256: str = "a" * 64, document_id: str = "doc_" + "b" * 64, page_number: int = 1) -> OutboundImageRecord:
    """按 adapter 的归一化结果生成 fake outbound 记录。"""
    prepared = analyzer.prepare_image(payload)
    return OutboundImageRecord(source_relative_path=source_relative_path, source_sha256=source_sha256, document_id=document_id, page_number=page_number, image_index=image_index, original_sha256=prepared.original_sha256, normalized_sha256=prepared.normalized_sha256, media_type=prepared.media_type, width=prepared.width, height=prepared.height, original_bytes=prepared.original_bytes, normalized_bytes=len(prepared.payload), transformation=prepared.transformation, context_sha256=sha256_bytes(context[:2000].encode("utf-8")), provider="wcode", model=analyzer.model, prompt_version=PROMPT_VERSION, response_adapter_version=RESPONSE_ADAPTER_VERSION, remote_policy_sha256=analyzer.remote_policy_sha256)


class MultimodalContractTests(unittest.TestCase):
    def test_candidate_identity_prefers_stable_unit_id_across_metadata_enrichment(self) -> None:
        """Dense metadata 增强不能在混合合并时重复 Chroma 单元。"""
        content = "unit content"
        sparse_metadata = {"unit_id": "unit_123", "page_number": 3}
        dense_metadata = {**sparse_metadata, "source": "source.pdf", "chunk_index": 7}
        self.assertEqual(candidate_identity(content, sparse_metadata), candidate_identity(content, dense_metadata))
        self.assertEqual(candidate_identity(content, sparse_metadata), "unit:unit_123")

    def test_hybrid_rerank_rewards_candidates_seen_by_both_retrievers(self) -> None:
        """混合证据在没有 gold 知识时也能获得共享检索器加分。"""
        metadata = {"unit_id": "unit_123", "modality": "text"}
        dense = {"page_content": "unit content", "metadata": metadata, "candidate_key": "unit:unit_123", "dense_score_norm": 0.8, "dense_score": 0.8, "sparse_score_norm": 0.0, "sparse_score": 0.0, "retrieval_sources": {"dense"}}
        sparse = {"page_content": "unit content", "metadata": metadata, "candidate_key": "unit:unit_123", "dense_score_norm": 0.0, "dense_score": 0.0, "sparse_score_norm": 0.4, "sparse_score": 0.4, "retrieval_sources": {"sparse"}}
        merged = _merge_candidates([dense], [sparse], return_before_final=True)
        self.assertEqual(len(merged), 1)
        self.assertAlmostEqual(merged[0]["rerank_score"], 0.65 * 0.8 + 0.25 * 0.4 + 0.2)

    def test_figure_word_does_not_create_a_modality_retrieval_branch(self) -> None:
        """用户提到图时仍走统一 dense 检索，不按内部 modality 改写结果。"""
        vector_db = MagicMock()
        vector_db.similarity_search_with_relevance_scores.return_value = [
            (MagicMock(page_content="图中内容", metadata={"unit_id": "unit_1", "modality": "image"}), 0.9),
        ]
        sparse_retriever = MagicMock()
        sparse_retriever.search.return_value = []

        trace = _build_retrieval_trace_with_resources(
            "图中是什么？",
            RagRetrievalConfig(),
            vector_db=vector_db,
            embedding_function=MagicMock(),
            sparse_retriever=sparse_retriever,
        )

        vector_db.similarity_search_with_relevance_scores.assert_called_once_with("图中是什么？", k=10)
        self.assertNotIn("dense_modality", {item["retrieval_source"] for item in trace["evidence_payload"]})

    """验证不依赖解析器、模型或远程 API 的核心安全契约。"""

    def test_bbox_rejects_reverse_coordinates(self) -> None:
        """边界框必须使用统一的左上至右下方向。"""
        with self.assertRaises(ValueError):
            BoundingBox(x0=0.8, y0=0.1, x1=0.2, y1=0.9)

    def test_retrieval_text_keeps_uncertain_relation_out_of_directed_relation(self) -> None:
        """不确定箭头只能作为不确定关系输出。"""
        analysis = VisionAnalysis(content_kind="causal_graph", ocr_text="", visible_facts=["存在连接线"], summary="", entities=[], table_markdown="", formula_latex="", directed_relations=[], uncertain_relations=["A ? B"], confidence=0.5, informative=True)
        text = render_retrieval_text(content_kind="causal_graph", analysis=analysis)
        self.assertIn("A ? B", text)
        self.assertNotIn("A -> B", text)

    def test_asset_store_rejects_escape_and_preserves_existing_asset(self) -> None:
        """资源存储不得允许目录逃逸或同 URI 覆盖。"""
        with tempfile.TemporaryDirectory() as directory:
            store = AssetStore(Path(directory))
            uri = store.put("doc_" + "a" * 64, "figure.png", b"one")
            self.assertEqual(store.read(uri), b"one")
            with self.assertRaises(ValueError):
                store.read("../secret")
            self.assertNotEqual(uri, store.put("doc_" + "a" * 64, "figure.png", b"two"))

    def test_asset_store_shortens_long_names_without_losing_content(self) -> None:
        """Windows 路径限制下，长资源名必须写入可读的确定性 URI。"""
        with tempfile.TemporaryDirectory() as directory:
            store = AssetStore(Path(directory) / ("nested_" * 12))
            name = "page_" + "a" * 180 + ".png"
            uri = store.put("doc_" + "a" * 64, name, b"image")
            self.assertLess(len(Path(uri).name), 80)
            self.assertEqual(store.read(uri), b"image")

    @patch("Agent.knowledge_base.multimodal.index._embeddings")
    @patch("Agent.knowledge_base.multimodal.index.Chroma")
    def test_staged_index_closes_client_after_write(self, chroma: MagicMock, _embeddings: MagicMock) -> None:
        """持久化写入在移动 staged 目录前释放 SQLite 客户端。"""
        db = MagicMock()
        db._collection.count.return_value = 1
        chroma.return_value = db
        unit = MagicMock(retrieval_text="causal text", unit_id="unit_1")
        unit.chroma_metadata.return_value = {}
        with tempfile.TemporaryDirectory() as directory:
            count = StagedIndex(Path(directory), "test_collection").write([unit])
        self.assertEqual(count, 1)
        db._client.close.assert_called_once_with()

    def test_inspect_is_deterministic_and_does_not_create_index(self) -> None:
        """inspect 只产生 manifest，不创建索引或资源目录。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "note.md"; source.write_text("因果图", encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "runtime" / "active.json")
            first = service.inspect([str(source)]); second = service.inspect([str(source)])
            self.assertEqual(first["manifest"], second["manifest"])
            self.assertEqual(first["manifest"]["parser"], "docling")
            self.assertIn("parser", first["configuration"])
            self.assertIn("docling_layout_model_available", first["configuration"]["parser"])
            self.assertFalse((root / "indexes").exists())
            self.assertFalse((root / "assets").exists())

    def test_csv_parser_keeps_header_and_row_semantics(self) -> None:
        """CSV 行的检索文本必须包含表头、行号和单元格值。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "data.csv"
            source.write_text("处理,结果\n1,改善\n", encoding="utf-8")
            parsed = parse_document(source, "mineru")
            self.assertEqual(parsed.parser_name, "csv")
            self.assertEqual(parsed.items[0].modality, "table")
            self.assertIn("行：2", parsed.items[0].raw_text)
            self.assertIn("处理：1", parsed.items[0].raw_text)
            self.assertIn("结果：改善", parsed.items[0].raw_text)

    def test_xlsx_parser_keeps_sheet_and_row_semantics(self) -> None:
        """XLSX 行必须包含工作表、行号和对应单元格语义。"""
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "data.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "实验"; sheet.append(["处理", "结果"]); sheet.append(["干预", "改善"]); workbook.save(source); workbook.close()
            parsed = parse_document(source, "mineru")
            self.assertEqual(parsed.parser_name, "xlsx")
            self.assertIn("工作表：实验", parsed.items[0].raw_text)
            self.assertIn("行：2", parsed.items[0].raw_text)

    def test_standalone_image_parser_does_not_run_local_rapidocr(self) -> None:
        """生产图片解析只保留资源，不再调用 RapidOCR 生成本地 OCR 文本。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.png"; source.write_bytes(_png_bytes())
            with patch("Agent.knowledge_base.multimodal.parsers._rapidocr_probe") as rapidocr:
                parsed = parse_document(source, "docling")
            rapidocr.assert_not_called()
            self.assertEqual(parsed.items[0].raw_text, "")

    def test_docling_pdf_parser_returns_page_scoped_text(self) -> None:
        """本地 Docling artifacts 必须能把 PDF 正文转换为一基页码单元。"""
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.pdf"
            pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, "Causal inference sample document"); pdf.save()
            with patch.dict(os.environ, {"MULTIMODAL_DOCLING_ARTIFACTS_DIR": str(REPOSITORY_ROOT / "Agent" / "knowledge_base" / "models" / "docling"), "MULTIMODAL_DOCLING_BATCH_SIZE": "1"}):
                parsed = parse_document(source, "mineru")
            self.assertEqual(parsed.parser_name, "docling")
            self.assertTrue(parsed.items, parsed.issues)
            self.assertTrue(any("Causal inference" in item.raw_text for item in parsed.items))
            self.assertTrue(any(item.page_number == 1 for item in parsed.items))

    def test_docling_pdf_parser_skips_blank_page_without_failing_source(self) -> None:
        """合法空白页应留下 warning，不能使整份 required PDF 失败。"""
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.pdf"
            pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, "Causal inference"); pdf.showPage(); pdf.showPage(); pdf.save()
            with patch.dict(os.environ, {"MULTIMODAL_DOCLING_ARTIFACTS_DIR": str(REPOSITORY_ROOT / "Agent" / "knowledge_base" / "models" / "docling"), "MULTIMODAL_DOCLING_BATCH_SIZE": "1"}):
                parsed = parse_document(source, "docling")
            self.assertTrue(parsed.items)
            self.assertFalse(any(issue.severity is IssueSeverity.ERROR for issue in parsed.issues))
            self.assertTrue(any(issue.code == "pdf_page_empty" for issue in parsed.issues))

    def test_docling_page_conversion_retries_once_with_new_converter(self) -> None:
        """converter 状态异常时应重建一次，并返回可供后续页复用的新实例。"""
        failed = MagicMock(); failed.convert.side_effect = RuntimeError("stale converter")
        replacement = MagicMock(); replacement.convert.return_value = "result"
        with patch("Agent.knowledge_base.multimodal.parsers._new_docling_converter", return_value=replacement) as factory:
            result, converter = _convert_docling_page(Path("sample.pdf"), 7, object(), failed)
        self.assertEqual(result, "result")
        self.assertIs(converter, replacement)
        factory.assert_called_once()
        replacement.convert.assert_called_once_with(Path("sample.pdf"), page_range=(7, 7))

    def test_docling_content_page_without_structure_uses_docling_page_image(self) -> None:
        """封面等整页图片未形成结构化项时仍须进入受审核的图片链。"""
        from Agent.knowledge_base.multimodal.parsers import _docling_page_items

        page_image = _png_bytes()
        with patch("Agent.knowledge_base.multimodal.parsers._docling_items", return_value=([], [])), patch("Agent.knowledge_base.multimodal.parsers._pdf_page_has_content", return_value=True), patch("Agent.knowledge_base.multimodal.parsers._docling_page_image_bytes", return_value=page_image):
            items, issues = _docling_page_items(MagicMock(), Path("cover.pdf"), 2)
        self.assertEqual(issues, ())
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].content_kind, "page_render")
        self.assertEqual(items[0].page_number, 2)
        self.assertEqual(items[0].asset_bytes, page_image)

    def test_evaluate_rejects_any_required_source_parse_error(self) -> None:
        """任一默认 required 来源解析失败时不得发布混合成功版本。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); version = "mm_" + "a" * 20; index_dir = root / "indexes" / version; index_dir.mkdir(parents=True)
            unit = KnowledgeUnit(unit_id="unit_" + "a" * 64, document_id="doc_" + "b" * 64, modality="text", content_kind="paragraph", raw_text="正文", retrieval_text="类型：paragraph\n正文", content_hash="c" * 64, parser_name="text", parser_version="1", embedding_provider="huggingface", embedding_model="model", status=UnitStatus.COMPLETED)
            (index_dir / "units.jsonl").write_text(unit.model_dump_json() + "\n", encoding="utf-8")
            issue = IngestionIssue(code="pdf_parse_failed", message="PDF 失败", severity=IssueSeverity.ERROR)
            (index_dir / "issues.jsonl").write_text(issue.model_dump_json() + "\n", encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "runtime" / "active.json")
            manifest = {"index_version": version, "embedding": __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()}
            (index_dir / "manifest.json").write_text(__import__("json").dumps(manifest), encoding="utf-8")
            with patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.count", return_value=1):
                result = service.evaluate(version)
            self.assertFalse(result["passed"])
            self.assertIn("required_source_failed", result["failures"])

    def test_evaluate_rejects_missing_original_or_parser_audit_assets(self) -> None:
        """发布门禁必须验证来源和解析器原始产物均可回读。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); version = "mm_" + "a" * 20; index_dir = root / "indexes" / version; index_dir.mkdir(parents=True)
            unit = KnowledgeUnit(unit_id="unit_" + "a" * 64, document_id="doc_" + "b" * 64, modality="text", content_kind="paragraph", raw_text="正文", retrieval_text="类型：paragraph\n正文", content_hash="c" * 64, parser_name="text", parser_version="1", embedding_provider="huggingface", embedding_model="model", status=UnitStatus.COMPLETED)
            (index_dir / "units.jsonl").write_text(unit.model_dump_json() + "\n", encoding="utf-8")
            (index_dir / "issues.jsonl").write_text("", encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "runtime" / "active.json")
            manifest = {"index_version": version, "embedding": __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint(), "sources": [], "documents": [{"document_id": unit.document_id, "content_hash": "c" * 64, "source_asset_uri": "missing/source.txt", "parser_artifact_uris": ["missing/parser.json"]}]}
            (index_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            with patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.count", return_value=1):
                result = service.evaluate(version)
            self.assertFalse(result["passed"])
            self.assertIn("missing_audit_asset", result["failures"])

    def test_evaluate_rejects_orphaned_standardized_unit(self) -> None:
        """标准化单元必须关联到 manifest 中经哈希验证的原始资料。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); version = "mm_" + "a" * 20; index_dir = root / "indexes" / version; index_dir.mkdir(parents=True)
            unit = KnowledgeUnit(unit_id="unit_" + "a" * 64, document_id="doc_" + "b" * 64, modality="text", content_kind="paragraph", raw_text="正文", retrieval_text="类型：paragraph\n正文", content_hash="c" * 64, parser_name="text", parser_version="1", embedding_provider="huggingface", embedding_model="model", status=UnitStatus.COMPLETED)
            (index_dir / "units.jsonl").write_text(unit.model_dump_json() + "\n", encoding="utf-8")
            (index_dir / "issues.jsonl").write_text("", encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "runtime" / "active.json")
            manifest = {"index_version": version, "embedding": __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint(), "sources": [], "documents": []}
            (index_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            with patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.count", return_value=1):
                result = service.evaluate(version)
            self.assertFalse(result["passed"])
            self.assertIn("missing_document_audit_chain", result["failures"])

    def test_directory_scan_preserves_relative_paths_for_same_names(self) -> None:
        """同一目录树中的同名文件必须保持不同来源路径。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source_root = root / "sources"
            for folder, text in (("one", "甲"), ("two", "乙")):
                folder_path = source_root / folder; folder_path.mkdir(parents=True)
                (folder_path / "same.md").write_text(text, encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "runtime" / "active.json")
            paths = [entry["relative_path"] for entry in service.inspect([str(source_root)])["manifest"]["sources"]]
            self.assertEqual(paths, ["sources/one/same.md", "sources/two/same.md"])

    def test_remote_vision_allowlist_rejects_arbitrary_local_file(self) -> None:
        """显式远程许可也不得扩大到非边界文件。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "private.png"; source.write_bytes(b"not-an-image")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "runtime" / "active.json")
            self.assertFalse(service._remote_resource_allowed(source, None))

    def test_remote_policy_allows_only_fixed_pearl_pages(self) -> None:
        """Pearl 资料必须同时匹配固定文件名和固定页码。"""
        service = MultimodalKnowledgeBaseMaintenance()
        pearl = Path(__file__).resolve().parents[1] / "Agent" / "knowledge_base" / "source" / "Pearl_2009_Causality-mono(1).pdf"
        self.assertTrue(service._remote_resource_allowed(pearl, 1))
        self.assertTrue(service._remote_resource_allowed(pearl, 7))
        self.assertTrue(service._remote_resource_allowed(pearl, 487))
        self.assertFalse(service._remote_resource_allowed(pearl, 0))
        self.assertFalse(service._remote_resource_allowed(pearl, 488))

    def test_vision_400_uses_json_prompt_fallback_and_audits_without_response(self) -> None:
        """结构化输出不兼容时只降级一次，并写入脱敏审计。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL}):
            image = _png_bytes(); context = "SENSITIVE_CONTEXT"
            response = MagicMock()
            response.choices = [MagicMock(message=MagicMock(content=_vision_payload(summary="RESPONSE_BODY")))]
            response.usage = MagicMock(prompt_tokens=3, completion_tokens=4)
            response._request_id = "request-1"
            client = MagicMock(); client.chat.completions.create.side_effect = [type("Error", (Exception,), {"status_code": 400})(), response]
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client):
                analysis = analyzer.analyze(image, "image/png", context, outbound_record=_outbound_record(analyzer, image, context))
            self.assertEqual(analysis.content_kind, "chart")
            self.assertEqual(client.chat.completions.create.call_count, 2)
            audit = (Path(directory) / "vision_audit.jsonl").read_text(encoding="utf-8")
            self.assertIn("success_json_prompt_fallback", audit)
            self.assertNotIn("content_kind", audit)
            for sensitive in (context, "RESPONSE_BODY", base64.b64encode(image).decode("ascii")[:32], "data:image"):
                self.assertNotIn(sensitive, audit)

    def test_vision_normalizes_rgb_and_enforces_outbound_hashes(self) -> None:
        """图片必须先真实解码、缩放、RGB 化，并与 outbound 记录逐字段匹配。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL, "VISION_MAX_PIXELS": "16"}):
            rgba = io.BytesIO(); Image.new("RGBA", (8, 8), (255, 0, 0, 80)).save(rgba, format="PNG"); image = rgba.getvalue()
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            prepared = analyzer.prepare_image(image)
            self.assertLessEqual(prepared.width * prepared.height, 16)
            self.assertEqual(prepared.media_type, "image/png")
            with Image.open(io.BytesIO(prepared.payload)) as decoded:
                self.assertEqual(decoded.mode, "RGB")
            with self.assertRaises(PermissionError):
                analyzer.analyze(image, "image/png")
            record = _outbound_record(analyzer, image).model_copy(update={"normalized_sha256": "0" * 64})
            with self.assertRaises(PermissionError):
                analyzer.analyze(image, "image/png", outbound_record=record)
            policy_drift = _outbound_record(analyzer, image).model_copy(update={"remote_policy_sha256": "0" * 64})
            with self.assertRaises(PermissionError):
                analyzer.analyze(image, "image/png", outbound_record=policy_drift)

    def test_vision_repairs_schema_once_and_audits_latency(self) -> None:
        """可提取但不符合 schema 的 JSON 只允许一次受控修复。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL}):
            image = _png_bytes(); broken = MagicMock(); repaired = MagicMock()
            broken.choices = [MagicMock(message=MagicMock(content='{"content_kind":"chart"}'))]
            repaired.choices = [MagicMock(message=MagicMock(content=_vision_payload()))]; repaired.usage = None; repaired._request_id = "repair"
            client = MagicMock(); client.chat.completions.create.side_effect = [broken, repaired]
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client):
                result = analyzer.analyze(image, "image/png", outbound_record=_outbound_record(analyzer, image))
            self.assertEqual(result.ocr_text, "图中文字")
            audit = json.loads((Path(directory) / "vision_audit.jsonl").read_text(encoding="utf-8").splitlines()[-1])
            self.assertEqual(audit["status"], "success_schema_repair")
            self.assertEqual(audit["retry_count"], 1)
            self.assertIsInstance(audit["latency_ms"], int)

    def test_vision_normalizes_text_lists_but_rejects_non_text_formula_values(self) -> None:
        """字符串列表应无损拼接；任意对象不能直接转成字符串。"""
        parsed = VisionAnalyzer._parse_analysis(_vision_payload(formula_latex=["x", "y"]))
        self.assertEqual(parsed.formula_latex, "x\ny")
        with self.assertRaisesRegex(ValueError, "invalid_schema"):
            VisionAnalyzer._parse_analysis(_vision_payload(formula_latex={"value": "x"}))

    def test_vision_schema_failure_records_only_sanitized_validation_paths(self) -> None:
        """Schema 失败应暴露字段/类型路径用于诊断，不能暴露响应正文。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL}):
            image = _png_bytes()
            broken = MagicMock(); broken.choices = [MagicMock(message=MagicMock(content=_vision_payload(confidence="not-a-number")))]
            repair = MagicMock(); repair.choices = [MagicMock(message=MagicMock(content=_vision_payload(confidence="not-a-number")))]
            client = MagicMock(); client.chat.completions.create.side_effect = [broken, repair]
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client):
                with self.assertRaisesRegex(RuntimeError, "without persisting"):
                    analyzer.analyze(image, "image/png", outbound_record=_outbound_record(analyzer, image))
            failure = next(Path(directory).glob("*.failure.json"))
            payload = json.loads(failure.read_text(encoding="utf-8"))
            self.assertEqual(payload["failure_category"], "invalid_schema")
            self.assertIn("confidence:float_parsing", payload["validation_error_paths"])
            self.assertNotIn("not-a-number", failure.read_text(encoding="utf-8"))

    def test_vision_retries_429_with_jitter_and_redacts_failure(self) -> None:
        """429 必须抖动退避；最终失败审计不得包含异常正文或密钥。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "super-secret", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL, "VISION_MAX_RETRIES": "1"}):
            image = _png_bytes(); response = MagicMock(); response.choices = [MagicMock(message=MagicMock(content=_vision_payload()))]; response.usage = None; response._request_id = "ok"
            rate_limit = type("RateLimit", (Exception,), {"status_code": 429})()
            client = MagicMock(); client.chat.completions.create.side_effect = [rate_limit, response]
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client), patch("Agent.knowledge_base.multimodal.vision.random.uniform", return_value=0.1), patch("Agent.knowledge_base.multimodal.vision.time.sleep") as sleep:
                analyzer.analyze(image, "image/png", outbound_record=_outbound_record(analyzer, image))
            sleep.assert_called_once_with(1.1)
            audit = (Path(directory) / "vision_audit.jsonl").read_text(encoding="utf-8")
            self.assertIn('"retry_count":1', audit)
            self.assertNotIn("super-secret", audit)

    def test_vision_failure_categories_cover_retry_and_deterministic_errors(self) -> None:
        """408/429/5xx/timeout/连接错误与确定性 4xx 必须得到稳定脱敏分类。"""
        cases = {408: "timeout", 429: "rate_limited", 500: "server_error", 503: "server_error", 401: "http_401", 403: "http_403", 404: "http_404"}
        for status, expected in cases.items():
            error = type(f"Http{status}", (Exception,), {"status_code": status})()
            self.assertEqual(VisionAnalyzer._failure_category(error), expected)
        self.assertEqual(VisionAnalyzer._failure_category(TimeoutError()), "timeout")
        self.assertEqual(VisionAnalyzer._failure_category(ConnectionError()), "connection")

    def test_vision_422_falls_back_and_budget_stops_second_image(self) -> None:
        """422 回退普通 JSON 一次；成功缓存外的第二张图片受预算阻断。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL}):
            first_image, second_image = _png_bytes("white"), _png_bytes("black")
            response = MagicMock(); response.choices = [MagicMock(message=MagicMock(content=_vision_payload()))]; response.usage = None; response._request_id = "fallback"
            unsupported = type("Unprocessable", (Exception,), {"status_code": 422})()
            client = MagicMock(); client.chat.completions.create.side_effect = [unsupported, response]
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client):
                analyzer.analyze(first_image, "image/png", outbound_record=_outbound_record(analyzer, first_image))
                with self.assertRaisesRegex(RuntimeError, "budget exhausted"):
                    analyzer.analyze(second_image, "image/png", outbound_record=_outbound_record(analyzer, second_image, image_index=2))
            self.assertEqual(client.chat.completions.create.call_count, 2)

    def test_vision_cache_prevents_second_sdk_call(self) -> None:
        """相同图片、模型和 Prompt 命中缓存后不得再次调用远程 SDK。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL}):
            image = _png_bytes(); response = MagicMock(); response.choices = [MagicMock(message=MagicMock(content=_vision_payload()))]; response.usage = None; response._request_id = "request"
            client = MagicMock(); client.chat.completions.create.return_value = response
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            outbound = _outbound_record(analyzer, image)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client):
                first = analyzer.analyze(image, "image/png", outbound_record=outbound)
                second = analyzer.analyze(image, "image/png", outbound_record=outbound)
            self.assertEqual(first, second)
            self.assertEqual(client.chat.completions.create.call_count, 1)

    def test_recorded_vision_failure_blocks_uncontrolled_retry(self) -> None:
        """失败状态必须在后续普通构建中阻止再次外发。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL, "VISION_MAX_RETRIES": "0"}):
            image = _png_bytes()
            client = MagicMock(); client.chat.completions.create.side_effect = RuntimeError("timeout")
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            outbound = _outbound_record(analyzer, image)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client):
                with self.assertRaises(RuntimeError): analyzer.analyze(image, "image/png", outbound_record=outbound)
            retry_client = MagicMock(); retry = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=1)
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=retry_client):
                with self.assertRaisesRegex(RuntimeError, "already recorded"): retry.analyze(image, "image/png", outbound_record=_outbound_record(retry, image))
            retry_client.assert_not_called()
            failure = next(Path(directory).glob("*.failure.json")).read_text(encoding="utf-8")
            self.assertIn("RuntimeError", failure); self.assertNotIn("timeout", failure)

    def test_unenriched_title_only_image_is_not_indexed(self) -> None:
        """未增强且无本地正文的独立图片不得作为低价值标题向量写入。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "image.png"; source.write_bytes(_png_bytes())
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            result = service.ingest([str(source)])
            manifest = json.loads((root / "indexes" / result["index_version"] / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(result["unit_count"], 0)
            self.assertEqual(manifest["quality_observations"]["low_value_images_skipped"], 1)

    def test_maintenance_run_is_idempotent_and_stops_before_publish_on_gate_failure(self) -> None:
        """维护编排复用既有版本，门禁失败或取消都不得切换 active pointer。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "source.md"; source.write_text("正文", encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            version = service._index_version(service._manifest(service._scan([str(source)])[0]))
            version_dir = root / "indexes" / version
            version_dir.mkdir(parents=True)
            for name in ("manifest.json", "units.jsonl", "issues.jsonl"):
                (version_dir / name).write_text("{}" if name.endswith(".json") else "", encoding="utf-8")
            (version_dir / "build_state.json").write_text(json.dumps({"status": "staged_complete", "unit_count": 0, "vector_count": 0}), encoding="utf-8")
            with patch.object(service, "ingest") as ingest, patch.object(service, "evaluate", return_value={"passed": False, "failures": ["quality_gate_failed"]}), patch.object(service, "publish") as publish:
                result = service.run([str(source)])
            self.assertEqual(result["status"], "gate_failed"); ingest.assert_not_called(); publish.assert_not_called()
            with patch.object(service, "publish") as publish:
                result = service.run([str(source)], cancel_check=lambda: True)
            self.assertEqual(result["status"], "cancelled"); publish.assert_not_called()

    def test_vision_rejects_model_or_provider_drift(self) -> None:
        """远程视觉边界必须固定为获准的 WCode 模型和域名。"""
        self.assertEqual(REQUIRED_MODEL, "qwen/qwen3-vl-8b-instruct")
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": "other/model"}):
            self.assertFalse(VisionAnalyzer(Path(directory), allow_remote_data=True).configured())
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.example.test/v1", "VISION_MODEL": REQUIRED_MODEL}):
            self.assertFalse(VisionAnalyzer(Path(directory), allow_remote_data=True).configured())
        for base_url in ("http://api.wcode.net/v1", "https://evilwcode.net/v1", "https://wcode.net.evil.test/v1"):
            with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": base_url, "VISION_MODEL": REQUIRED_MODEL}):
                self.assertFalse(VisionAnalyzer(Path(directory), allow_remote_data=True).configured())

    def test_vision_concurrency_cap_serializes_extra_requests(self) -> None:
        """并发上限必须约束实际 SDK 调用，而不是只保留配置字段。"""
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"VISION_API_KEY": "key", "VISION_BASE_URL": "https://api.wcode.net/v1", "VISION_MODEL": REQUIRED_MODEL, "VISION_MAX_CONCURRENCY": "1"}):
            response = MagicMock(); response.choices = [MagicMock(message=MagicMock(content=_vision_payload()))]; response.usage = None; response._request_id = "request"
            active = 0; peak = 0; lock = threading.Lock()
            def invoke(**kwargs):
                """模拟可并发观测的单次远程 SDK 调用。"""
                nonlocal active, peak
                with lock:
                    active += 1; peak = max(peak, active)
                time.sleep(0.03)
                with lock: active -= 1
                return response
            client = MagicMock(); client.chat.completions.create.side_effect = invoke
            analyzer = VisionAnalyzer(Path(directory), allow_remote_data=True, max_images=2)
            images = [_png_bytes("white"), _png_bytes("black")]
            with patch("Agent.knowledge_base.multimodal.vision.OpenAI", return_value=client):
                workers = [threading.Thread(target=analyzer.analyze, args=(image, "image/png"), kwargs={"outbound_record": _outbound_record(analyzer, image, image_index=index + 1)}) for index, image in enumerate(images)]
                for worker in workers: worker.start()
                for worker in workers: worker.join()
            self.assertEqual(peak, 1)

    def test_pipeline_remote_context_is_same_page_and_capped(self) -> None:
        """远程请求上下文只取当前图片 caption 与同页正文，最长 2000 字。"""
        image = ParsedItem("image", "image", raw_text="图片标题", page_number=2, asset_bytes=_png_bytes())
        items = [
            image,
            ParsedItem("text", "paragraph", raw_text="同页正文" * 400, page_number=2),
            ParsedItem("text", "paragraph", raw_text="其他页正文", page_number=3),
            ParsedItem("image", "image", raw_text="其他图片", page_number=2, asset_bytes=_png_bytes("black")),
        ]
        context = MultimodalKnowledgeBaseMaintenance._same_page_context(image, items)
        self.assertTrue(context.startswith("图片标题"))
        self.assertLessEqual(len(context), 2000)
        self.assertNotIn("其他页正文", context)
        self.assertNotIn("其他图片", context)

    def test_only_complete_staged_version_can_be_reused(self) -> None:
        """崩溃留下的确定性版本目录不得被 run 当作已完成索引复用。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            version = "mm_" + "a" * 20
            version_dir = root / "indexes" / version
            version_dir.mkdir(parents=True)
            service = MultimodalKnowledgeBaseMaintenance(
                asset_root=root / "assets",
                index_root=root / "indexes",
                active_config=root / "active.json",
            )

            self.assertFalse(service._is_reusable_staged_version(version))
            (version_dir / "build_state.json").write_text(
                json.dumps({"status": "staged_complete", "unit_count": 2, "vector_count": 1}),
                encoding="utf-8",
            )
            self.assertFalse(service._is_reusable_staged_version(version))
            (version_dir / "build_state.json").write_text(
                json.dumps({"status": "staged_complete", "unit_count": 0, "vector_count": 0}),
                encoding="utf-8",
            )
            for name in ("manifest.json", "units.jsonl", "issues.jsonl"):
                (version_dir / name).write_text("{}" if name.endswith(".json") else "", encoding="utf-8")
            self.assertTrue(service._is_reusable_staged_version(version))

    def test_failed_vector_build_resumes_from_page_checkpoints(self) -> None:
        """页解析完成后向量写入失败，重试不得再次运行已完成页面。"""
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "two-pages.pdf"
            pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, "page one"); pdf.showPage(); pdf.drawString(72, 720, "page two"); pdf.save()
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            parsed_pages = [
                ParsedDocument("docling", "2.115.0", (ParsedItem("text", "paragraph", raw_text="page one has enough searchable causal inference content", page_number=1),), raw_artifacts=(("docling_page_0001.json", b"{}"),)),
                ParsedDocument("docling", "2.115.0", (ParsedItem("text", "paragraph", raw_text="page two has enough searchable causal inference content", page_number=2),), raw_artifacts=(("docling_page_0002.json", b"{}"),)),
            ]
            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", side_effect=parsed_pages) as parse_page, patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.write", side_effect=RuntimeError("vector failed")):
                with self.assertRaisesRegex(RuntimeError, "vector failed"):
                    service.ingest([str(source)])
            self.assertEqual(parse_page.call_count, 2)

            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page") as parse_page, patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.write", return_value=2):
                result = service.ingest([str(source)])
            parse_page.assert_not_called()
            self.assertEqual(result["unit_count"], 2)
            self.assertEqual(result["vector_count"], 2)

    def test_retry_candidate_reuses_only_non_error_page_checkpoints(self) -> None:
        """新的 retry 候选应只重新解析带 ERROR issue 的页，不能修改旧版本。"""
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "three-pages.pdf"
            pdf = canvas.Canvas(str(source))
            for page in ("one", "two", "three"):
                pdf.drawString(72, 720, page); pdf.showPage()
            pdf.save()
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            failed_issue = IngestionIssue(code="pdf_page_parse_failed", message="transient", severity=IssueSeverity.ERROR, blocking=True, source_path=str(source))
            failed = ParsedDocument("docling", "2.115.0", (), issues=(failed_issue,), raw_artifacts=(("docling_page_0001.json", b"{}"),))
            stable_two = ParsedDocument("docling", "2.115.0", (ParsedItem("text", "paragraph", raw_text="page two has enough searchable causal inference content", page_number=2),), raw_artifacts=(("docling_page_0002.json", b"{}"),))
            stable_three = ParsedDocument("docling", "2.115.0", (ParsedItem("text", "paragraph", raw_text="page three has enough searchable causal inference content", page_number=3),), raw_artifacts=(("docling_page_0003.json", b"{}"),))
            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", side_effect=[failed, stable_two, stable_three]), patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.write", side_effect=[2, 3]):
                original = service.ingest([str(source)])
                original_manifest = json.loads((root / "indexes" / original["index_version"] / "manifest.json").read_text(encoding="utf-8"))
                repaired = ParsedDocument("docling", "2.115.0", (ParsedItem("text", "paragraph", raw_text="page one now parses after resources recovered", page_number=1),), raw_artifacts=(("docling_page_0001.json", b"{}"),))
                with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=repaired) as parse_page:
                    retry = service.ingest([str(source)], retry_failed=True, retry_generation=1, retry_from_index_version=original["index_version"])
            parse_page.assert_called_once_with(source, "docling", 1)
            self.assertNotEqual(retry["index_version"], original["index_version"])
            self.assertEqual(retry["unit_count"], 3)
            self.assertEqual(original_manifest["index_version"], original["index_version"])

    def test_index_version_changes_with_embedding_or_vision_configuration(self) -> None:
        """影响向量或增强产物的配置变化必须生成新不可变版本。"""
        with tempfile.TemporaryDirectory() as directory:
            service = MultimodalKnowledgeBaseMaintenance(asset_root=Path(directory) / "assets", index_root=Path(directory) / "indexes", active_config=Path(directory) / "active.json")
            entries = [{"path": "unused", "relative_path": "source.md", "content_hash": "a" * 64}]
            baseline = service._index_version(service._manifest(entries, allow_remote_data=False))
            with patch.dict(os.environ, {"VISION_MODEL": "changed"}):
                changed_vision = service._index_version(service._manifest(entries, allow_remote_data=False))
            with patch("Agent.knowledge_base.multimodal.pipeline.embedding_fingerprint", return_value={"provider": "openai_compatible", "model": "changed", "mode": "api", "normalized": False}):
                changed_embedding = service._index_version(service._manifest(entries, allow_remote_data=False))
            self.assertNotEqual(baseline, changed_vision); self.assertNotEqual(baseline, changed_embedding)

    def test_quality_gate_blocks_complete_but_insufficient_enrichment(self) -> None:
        """质量不足即使完整性通过也必须阻止发布。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); version = "mm_" + "a" * 20; index_dir = root / "indexes" / version; index_dir.mkdir(parents=True)
            embedding = __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "runtime" / "active.json")
            content = b"source"; source_hash = sha256_bytes(content); document_id = stable_id("doc", {"path": "source.md", "content_hash": source_hash})
            source_uri = AssetStore(root / "assets").put(document_id, "source.md", content, category="source")
            unit = KnowledgeUnit(unit_id="unit_" + "a" * 64, document_id=document_id, modality="text", content_kind="paragraph", raw_text="正文", retrieval_text="正文", content_hash=sha256_bytes("正文".encode()), parser_name="text", parser_version="1", embedding_provider=embedding["provider"], embedding_model=embedding["model"], status=UnitStatus.COMPLETED)
            manifest = {"index_version": version, "embedding": embedding, "build_configuration": {}, "sources": [{"relative_path": "source.md", "content_hash": source_hash}], "documents": [{"document_id": document_id, "relative_path": "source.md", "content_hash": source_hash, "source_asset_uri": source_uri, "parser_artifacts": []}], "quality_policy": {"min_eligible_images": 1, "min_enriched_images": 1, "min_enrichment_rate": 1.0, "max_vision_failed_images": 0}, "quality_observations": {"eligible_images": 1, "enriched_images": 0, "vision_failed_images": 0, "skipped_images": 1}}
            (index_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8"); (index_dir / "units.jsonl").write_text(unit.model_dump_json() + "\n", encoding="utf-8"); (index_dir / "issues.jsonl").write_text("", encoding="utf-8")
            with patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.count", return_value=1):
                result = service.evaluate(version)
            self.assertFalse(result["passed"]); self.assertTrue(result["quality"]["failures"]); self.assertEqual(result["failures"], ["quality_gate_failed"])
            with self.assertRaisesRegex(ValueError, "blocking evaluation"):
                service.publish(version)
            manifest["quality_observations"]["enriched_images"] = 1
            (index_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            with patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.count", return_value=1):
                result = service.evaluate(version)
            self.assertTrue(result["passed"]); self.assertEqual(service.publish(version)["status"], "published")

    def test_legacy_manifest_cannot_pass_new_quality_or_publication_gates(self) -> None:
        """缺少 P0 构建和质量契约的历史版本不得被默认零值放行。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); version = "mm_" + "a" * 20; index_dir = root / "indexes" / version; index_dir.mkdir(parents=True)
            embedding = __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()
            unit = KnowledgeUnit(unit_id="unit_" + "a" * 64, document_id="doc_" + "b" * 64, modality="text", content_kind="paragraph", raw_text="正文", retrieval_text="正文", content_hash=sha256_bytes("正文".encode()), parser_name="text", parser_version="1", embedding_provider=embedding["provider"], embedding_model=embedding["model"], status=UnitStatus.COMPLETED)
            (index_dir / "manifest.json").write_text(json.dumps({"index_version": version, "embedding": embedding, "sources": [], "documents": []}), encoding="utf-8")
            (index_dir / "units.jsonl").write_text(unit.model_dump_json() + "\n", encoding="utf-8"); (index_dir / "issues.jsonl").write_text("", encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            with patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.count", return_value=1):
                result = service.evaluate(version)
            self.assertFalse(result["passed"]); self.assertIn("legacy_manifest_missing_build_configuration", result["failures"]); self.assertIn("legacy_manifest_missing_quality_policy", result["quality"]["failures"])
            (index_dir / "evaluation.json").write_text(json.dumps({"passed": True}), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "legacy manifest"):
                service.publish(version)

    def test_multimodal_is_the_default_and_only_rag_corpus(self) -> None:
        """未声明 corpus 时必须默认多模态，旧 medical 输入必须显式失败。"""
        payload = {"questions": [{"question": "解释文档", "intent": "证据", "priority": "high", "why_needed": "需要页面"}]}
        questions = normalize_rag_question_output(payload, 3)
        self.assertEqual(questions[0]["corpus"], "multimodal")

        payload["questions"][0]["corpus"] = "medical"
        with self.assertRaises(ValueError):
            normalize_rag_question_output(payload, 3)

    def test_default_rag_registry_keeps_original_tool_name(self) -> None:
        """默认 RAG 继续使用原工具名，但底层只绑定多模态 Service。"""
        self.assertEqual([tool.name for tool in build_rag_tools()], ["rag_enrichment_search"])

    def test_rag_runtime_defaults_to_published_multimodal_index(self) -> None:
        """原 RagRuntime 必须从多模态 active pointer 解析默认 collection。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            version = "mm_test"
            version_dir = root / "indexes" / version
            (version_dir / "chroma").mkdir(parents=True)
            manifest = {"index_version": version, "embedding": {"provider": "local", "model": "model", "mode": "local", "normalized": True}}
            manifest_path = version_dir / "manifest.json"
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            active_path = root / "runtime" / "active.json"
            active_path.parent.mkdir()
            active_path.write_text(json.dumps({
                "index_version": version,
                "index_path": f"{version}/chroma",
                "collection_name": "causal_multimodal_mm_test",
                "manifest_sha256": __import__("hashlib").sha256(manifest_path.read_bytes()).hexdigest(),
                "embedding": manifest["embedding"],
            }), encoding="utf-8")
            embedding_config = {"status": "ready", "mode": "local", "provider": "local", "model": "model", "path": "model"}
            with patch.dict(os.environ, {"MULTIMODAL_INDEX_ROOT": str(root / "indexes"), "MULTIMODAL_ACTIVE_INDEX_CONFIG": str(active_path), "MULTIMODAL_ALLOW_NON_PRODUCTION_ACTIVE": "true"}), patch("Agent.knowledge_base.rag_runtime.resolve_embedding_runtime_config", return_value=embedding_config):
                config = RagRuntimeConfig.from_environment()
            self.assertEqual(Path(config.vector_db_dir), version_dir / "chroma")
            self.assertEqual(config.collection_name, "causal_multimodal_mm_test")
            self.assertEqual(config.release_id, version)

    def test_original_rag_metadata_normalizer_preserves_multimodal_identity(self) -> None:
        """原检索链必须识别多模态 document、page、asset 和内容类型字段。"""
        from Agent.knowledge_base.query_rag import _normalize_chunk_metadata

        metadata = _normalize_chunk_metadata({
            "document_id": "doc_1",
            "page_number": 3,
            "asset_uri": "doc_1/images/chart.png",
            "modality": "image",
            "content_kind": "chart",
            "content_hash": "abc",
        }, "图表检索文本")
        self.assertEqual(metadata["doc_id"], "doc_1")
        self.assertEqual(metadata["page"], 3)
        self.assertEqual(metadata["source_name"], "chart.png")
        self.assertEqual(metadata["doc_type"], "chart")
        self.assertEqual(metadata["corpus"], "multimodal")

    def test_rapidocr_probe_classifies_engine_unavailable_failed_and_no_text(self) -> None:
        """RapidOCR 必须区分依赖缺失、推理失败和无文字三种状态，不能统一返回空串。"""
        fp_available = {"engine": "rapidocr-onnxruntime", "engine_version": "1.4.4", "model_fingerprint": "1.4.4", "status": "available"}
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value={**fp_available, "status": "unavailable"}):
            self.assertEqual(_rapidocr_probe(b"image").status, "engine_unavailable")
        broken_engine = MagicMock(); broken_engine.side_effect = RuntimeError("inference failed")
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value=fp_available), patch("Agent.knowledge_base.multimodal.parsers._RAPID_OCR_ENGINE", broken_engine):
            probe = _rapidocr_probe(b"image")
            self.assertEqual(probe.status, "failed")
            self.assertEqual(probe.failure_type, "RuntimeError")
        empty_result = MagicMock(); empty_result.txts = []; empty_result.scores = []
        empty_engine = MagicMock(); empty_engine.return_value = empty_result
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value=fp_available), patch("Agent.knowledge_base.multimodal.parsers._RAPID_OCR_ENGINE", empty_engine):
            self.assertEqual(_rapidocr_probe(b"image").status, "no_text")

    def test_rapidocr_initialization_failure_is_engine_unavailable(self) -> None:
        """模型或引擎初始化失败必须阻断，不能伪装成单图推理失败。"""
        fp_available = {"engine": "rapidocr-onnxruntime", "engine_version": "1.4.4", "model_fingerprint": "a" * 64, "status": "available"}
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value=fp_available), patch("Agent.knowledge_base.multimodal.parsers._RAPID_OCR_ENGINE", None), patch("rapidocr.RapidOCR", side_effect=FileNotFoundError("model missing")):
            probe = _rapidocr_probe(b"image")
        self.assertEqual(probe.status, "engine_unavailable")
        self.assertEqual(probe.failure_type, "FileNotFoundError")

    def test_ocr_engine_unavailable_produces_blocking_issue(self) -> None:
        """依赖缺失必须生成 ERROR 级 issue，使正式来源在门禁处被阻断。"""
        from Agent.knowledge_base.multimodal.parsers import _ocr_issues, OcrResult
        probe = OcrResult("", "rapidocr-onnxruntime", "1.4.4", "1.4.4", 0, 0.0, 0, "engine_unavailable")
        issues = _ocr_issues(probe, Path("source.pdf"))
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].code, "ocr_engine_unavailable")
        self.assertEqual(issues[0].severity, IssueSeverity.ERROR)
        self.assertTrue(issues[0].blocking)

    def test_informative_false_keeps_remote_ocr_unit(self) -> None:
        """语义不具信息量但远程 OCR 有文本时仍生成一个完整图片单元。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); image = _png_bytes(); source = root / "page.png"; source.write_bytes(image)
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            embedding = __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()
            store = AssetStore(root / "assets")
            preparer = VisionAnalyzer(root / "cache", allow_remote_data=False)
            analyzer = MagicMock(); analyzer.model = REQUIRED_MODEL; analyzer.response_adapter_version = RESPONSE_ADAPTER_VERSION; analyzer.remote_policy_sha256 = preparer.remote_policy_sha256
            analyzer.analyze.return_value = VisionAnalysis(content_kind="illustration", ocr_text="远程 OCR 文本", visible_facts=[], summary="", entities=[], table_markdown="", formula_latex="", directed_relations=[], uncertain_relations=[], informative=False, confidence=0.3)
            quality = {"eligible_images": 0, "enriched_images": 0, "vision_failed_images": 0, "skipped_images": 0, "low_value_images_skipped": 0, "filtered_short_text_units": 0}
            issues: list = []
            item = ParsedItem(modality="image", content_kind="image", raw_text="本地 caption", asset_bytes=image, asset_name="page_0001_image_0001.png", page_number=1)
            outbound = [_outbound_record(preparer, image, source_relative_path="page.png", source_sha256="b" * 64, document_id="doc_" + "a" * 64)]
            with patch.object(service, "_remote_resource_allowed", return_value=True):
                unit = service._build_unit(item, source, "doc_" + "a" * 64, 1, "docling", "2.115.0", embedding, store, analyzer, quality, issues, allow_remote_data=True, source_relative_path="page.png", source_sha256="b" * 64, outbound_records=outbound)
            self.assertIsNotNone(unit)
            self.assertIn("远程 OCR 文本", unit.retrieval_text)
            self.assertIn("本地 caption", unit.raw_text)
            self.assertEqual(unit.vision_model, REQUIRED_MODEL)
            self.assertEqual(len(outbound), 1)
            manifest = outbound[0].model_dump(mode="json")
            self.assertEqual(manifest["source_sha256"], "b" * 64)
            self.assertEqual(manifest["page_number"], 1)
            self.assertEqual(manifest["original_sha256"], sha256_bytes(image))
            self.assertNotIn("本地 caption", json.dumps(manifest, ensure_ascii=False))

    def test_remote_failure_blocks_image_unit(self) -> None:
        """统一远程分析失败时不得降级为本地 OCR 或 caption 单元。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); image = _png_bytes(); source = root / "page.png"; source.write_bytes(image)
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            embedding = __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()
            store = AssetStore(root / "assets")
            preparer = VisionAnalyzer(root / "cache", allow_remote_data=False)
            analyzer = MagicMock(); analyzer.model = REQUIRED_MODEL; analyzer.response_adapter_version = RESPONSE_ADAPTER_VERSION; analyzer.remote_policy_sha256 = preparer.remote_policy_sha256
            analyzer.analyze.side_effect = RuntimeError("timeout")
            quality = {"eligible_images": 0, "enriched_images": 0, "vision_failed_images": 0, "skipped_images": 0, "low_value_images_skipped": 0, "filtered_short_text_units": 0}
            issues: list = []
            item = ParsedItem(modality="image", content_kind="image", raw_text="本地 caption", asset_bytes=image, asset_name="page_0001_image_0001.png", page_number=1)
            outbound = [_outbound_record(preparer, image, source_relative_path="page.png", source_sha256="b" * 64, document_id="doc_" + "a" * 64)]
            with patch.object(service, "_remote_resource_allowed", return_value=True):
                unit = service._build_unit(item, source, "doc_" + "a" * 64, 1, "docling", "2.115.0", embedding, store, analyzer, quality, issues, allow_remote_data=True, source_relative_path="page.png", source_sha256="b" * 64, outbound_records=outbound)
            self.assertIsNone(unit)
            self.assertTrue(any(issue.code == "remote_image_failed" and issue.blocking for issue in issues))

    def test_image_content_hash_uses_asset_bytes(self) -> None:
        """图片单元的 content_hash 必须基于图片字节而非 OCR 文本。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "page.png"; source.write_bytes(b"image-bytes")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            embedding = __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()
            store = AssetStore(root / "assets")
            analyzer = MagicMock(); analyzer.model = ""
            quality = {"eligible_images": 0, "enriched_images": 0, "vision_failed_images": 0, "skipped_images": 0, "low_value_images_skipped": 0, "filtered_short_text_units": 0}
            issues: list = []
            item = ParsedItem(modality="image", content_kind="image", raw_text="OCR text", asset_bytes=b"image-bytes", asset_name="page.png", page_number=1)
            unit = service._build_unit(item, source, "doc_" + "a" * 64, 1, "docling", "2.115.0", embedding, store, analyzer, quality, issues, allow_remote_data=False)
            self.assertEqual(unit.content_hash, sha256_bytes(b"image-bytes"))
            self.assertNotEqual(unit.content_hash, sha256_bytes("OCR text".encode()))

    def test_manifest_disables_local_ocr_and_versions_remote_limits(self) -> None:
        """manifest 必须移除生产 RapidOCR 依赖并版本化远程边界。"""
        service = MultimodalKnowledgeBaseMaintenance()
        manifest = service._manifest([], allow_remote_data=False)
        build_config = manifest["build_configuration"]
        self.assertNotIn("ocr", build_config)
        self.assertFalse(build_config["vision"]["local_ocr_enabled"])
        self.assertEqual(build_config["vision"]["prompt_version"], "vision-v2")
        self.assertGreater(build_config["vision"]["max_pixels"], 0)
        self.assertGreater(build_config["vision"]["max_image_bytes"], 0)
        self.assertIsNone(build_config["vision"]["max_images"])
        self.assertEqual(VisionAnalyzer(Path(tempfile.gettempdir()), allow_remote_data=False, max_images=101).max_images, 101)
        self.assertIn("remote_policy_hash", build_config)
        self.assertRegex(build_config["remote_policy_hash"], r"^[0-9a-f]{64}$")

    def test_outbound_manifest_audit_detects_tampering(self) -> None:
        """候选门禁必须回读 outbound 文件并校验哈希、来源与唯一图片身份。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); analyzer = VisionAnalyzer(root / "cache", allow_remote_data=False); image = _png_bytes()
            record = _outbound_record(analyzer, image)
            path = root / "outbound_manifest.json"
            MultimodalKnowledgeBaseMaintenance._write_outbound_manifest(path, [record])
            manifest = {"sources": [{"relative_path": "source.pdf", "content_hash": "a" * 64}], "build_configuration": {"remote_policy_hash": analyzer.remote_policy_sha256, "vision": {"response_adapter_version": RESPONSE_ADAPTER_VERSION}}, "outbound_manifest_sha256": sha256_bytes(path.read_bytes()), "outbound_image_count": 1}
            self.assertEqual(MultimodalKnowledgeBaseMaintenance._audit_outbound_manifest(root, manifest), [])
            path.write_text("{}", encoding="utf-8")
            self.assertEqual(MultimodalKnowledgeBaseMaintenance._audit_outbound_manifest(root, manifest), ["outbound_manifest_mismatch"])

    def test_frozen_outbound_manifest_requires_current_sources_and_policy(self) -> None:
        """远程运行只能使用预冻结、同来源且同策略的批准图片清单。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            analyzer = VisionAnalyzer(root / "cache", allow_remote_data=False, remote_policy_sha256=service.remote_policy.policy_sha256)
            record = _outbound_record(analyzer, _png_bytes())
            path = root / "approved.json"
            service._write_outbound_manifest(path, [record])
            entries = [{"relative_path": "source.pdf", "content_hash": "a" * 64}]
            self.assertEqual(service._frozen_outbound_records(path, entries), [record])
            self.assertEqual(service._frozen_outbound_records(None, entries), [])
            with self.assertRaisesRegex(ValueError, "current frozen sources"):
                service._frozen_outbound_records(path, [{"relative_path": "source.pdf", "content_hash": "c" * 64}])

    def test_prepare_outbound_manifest_is_local_and_consumable_by_r2(self) -> None:
        """本地准备命令生成的记录应可直接作为 R2 的冻结输入，不调用远程 adapter。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "page.png"; image = _png_bytes(); source.write_bytes(image)
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            parsed = ParsedDocument("builtin-2", "2", (
                ParsedItem("text", "paragraph", raw_text="同页正文", page_number=1),
                ParsedItem("image", "image", raw_text="图片标题", page_number=1, asset_bytes=image, asset_name="page.png"),
            ))
            output = root / "approved.json"
            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=parsed), patch.object(service, "_remote_resource_allowed", return_value=True):
                result = service.prepare_outbound_manifest([str(source)], output)
            self.assertEqual(result["status"], "prepared")
            self.assertEqual(result["outbound_image_count"], 1)
            self.assertEqual(result["prepared_page_count"], 1)
            entries, _ = service._scan([str(source)])
            records = service._frozen_outbound_records(output, entries)
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].model, REQUIRED_MODEL)
            self.assertEqual(records[0].context_sha256, sha256_bytes("同页正文\n图片标题".encode("utf-8")))

    def test_prepare_outbound_manifest_stops_before_remote_on_parse_error(self) -> None:
        """本地解析失败时不得产出可被误用的远程外发清单。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "page.png"; source.write_bytes(_png_bytes())
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            failed = ParsedDocument("docling", "2.115.0", issues=(IngestionIssue(code="pdf_page_parse_failed", message="failed", severity=IssueSeverity.ERROR, blocking=True),))
            output = root / "approved.json"
            with patch.object(service, "_remote_resource_allowed", return_value=True), patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=failed):
                with self.assertRaisesRegex(ValueError, "local parsing failed"):
                    service.prepare_outbound_manifest([str(source)], output)
            self.assertFalse(output.exists())

    def test_prepare_outbound_manifest_limits_eligible_pages_before_parsing(self) -> None:
        """R2 smoke 的页数上限应在本地解析前生效，避免为找图片扫描完整 PDF。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "page.png"; source.write_bytes(_png_bytes())
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            parsed = ParsedDocument("docling", "2.115.0", ())
            output = root / "approved.json"
            with patch.object(service, "_source_page_count", return_value=20), patch.object(service, "_remote_resource_allowed", return_value=True), patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=parsed) as parse_page:
                result = service.prepare_outbound_manifest([str(source)], output, max_pages=2)
            self.assertEqual(result["prepared_page_count"], 2)
            self.assertEqual(parse_page.call_count, 2)

    def test_remote_image_without_frozen_record_is_blocked_before_sdk_call(self) -> None:
        """候选内不存在批准记录时，图片不得自行补写清单或调用远程 adapter。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); image = _png_bytes(); source = root / "page.png"; source.write_bytes(image)
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            embedding = __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()
            analyzer = MagicMock(); analyzer.model = REQUIRED_MODEL; analyzer.response_adapter_version = RESPONSE_ADAPTER_VERSION; analyzer.remote_policy_sha256 = service.remote_policy.policy_sha256
            quality = {"eligible_images": 0, "enriched_images": 0, "vision_failed_images": 0, "skipped_images": 0, "low_value_images_skipped": 0, "filtered_short_text_units": 0}
            issues: list = []
            item = ParsedItem(modality="image", content_kind="image", raw_text="caption", asset_bytes=image, asset_name="page.png", page_number=1)
            with patch.object(service, "_remote_resource_allowed", return_value=True):
                unit = service._build_unit(item, source, "doc_" + "a" * 64, 1, "docling", "2.115.0", embedding, AssetStore(root / "assets"), analyzer, quality, issues, allow_remote_data=True, source_relative_path="page.png", source_sha256="b" * 64, outbound_records=[])
            self.assertIsNone(unit)
            analyzer.analyze.assert_not_called()
            self.assertTrue(any(issue.code == "remote_image_failed" and issue.blocking for issue in issues))

    def test_image_unit_id_changes_with_prompt_fingerprint(self) -> None:
        """图片 unit ID 不再绑定 RapidOCR，但必须绑定远程提示词版本。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "page.png"; source.write_bytes(b"image-bytes")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            embedding = __import__("Agent.knowledge_base.multimodal.index", fromlist=["embedding_fingerprint"]).embedding_fingerprint()
            item = ParsedItem(modality="image", content_kind="image", raw_text="OCR text", asset_bytes=b"image-bytes", asset_name="page.png", page_number=1)
            analyzer = MagicMock(); analyzer.model = ""
            quality = {"eligible_images": 0, "enriched_images": 0, "vision_failed_images": 0, "skipped_images": 0, "low_value_images_skipped": 0, "filtered_short_text_units": 0}
            with patch("Agent.knowledge_base.multimodal.pipeline.PROMPT_VERSION", "vision-v1"):
                first = service._build_unit(item, source, "doc_" + "a" * 64, 1, "docling", "2.115.0", embedding, AssetStore(root / "assets"), analyzer, quality, [], allow_remote_data=False)
            with patch("Agent.knowledge_base.multimodal.pipeline.PROMPT_VERSION", "vision-v2"):
                second = service._build_unit(item, source, "doc_" + "a" * 64, 1, "docling", "2.115.0", embedding, AssetStore(root / "assets"), analyzer, quality, [], allow_remote_data=False)
            self.assertNotEqual(first.unit_id, second.unit_id)

    def test_local_parse_checkpoint_reuse_skips_docling_across_vision_policy(self) -> None:
        """--reuse-local-checkpoints-from 必须只复用本地解析，跨视觉策略跳过 Docling。"""
        from reportlab.pdfgen import canvas
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "page.pdf"
            pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, "causal inference content"); pdf.save()
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            parsed = ParsedDocument("docling", "2.115.0", (ParsedItem("text", "paragraph", raw_text="causal inference content", page_number=1),), raw_artifacts=(("docling_page_0001.json", b"{}"),))
            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=parsed), patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.write", return_value=1):
                ocr_only = service.ingest([str(source)])
            checkpoint_database = root / "indexes" / ocr_only["index_version"] / "checkpoints.sqlite3"
            self.assertTrue(checkpoint_database.is_file())
            connection = sqlite3.connect(checkpoint_database)
            try:
                rows = connection.execute("SELECT checkpoint_json FROM local_parse_checkpoints").fetchall()
            finally:
                connection.close()
            self.assertEqual(len(rows), 1)
            self.assertEqual(json.loads(rows[0][0])["schema_version"], "local-parse-v2")
            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page") as parse_page, patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.write", return_value=1):
                vlm_version = service.ingest([str(source)], allow_remote_data=True, max_images=1, reuse_local_from_index_version=ocr_only["index_version"])
            parse_page.assert_not_called()
            self.assertNotEqual(vlm_version["index_version"], ocr_only["index_version"])

    def test_reuse_local_and_retry_are_mutually_exclusive(self) -> None:
        """--reuse-local-checkpoints-from 和 --retry-from-index-version 不得同时使用。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "s.md"; source.write_text("正文", encoding="utf-8")
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            with self.assertRaisesRegex(ValueError, "cannot be combined"):
                service.ingest([str(source)], retry_failed=True, retry_from_index_version="mm_" + "a" * 20, reuse_local_from_index_version="mm_" + "b" * 20)

    def test_invalid_local_checkpoint_reparses_instead_of_silently_reusing(self) -> None:
        """schema、资源或 parser 契约失配时必须回退本地解析。"""
        from reportlab.pdfgen import canvas
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / "page.pdf"
            pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, "causal inference content"); pdf.save()
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            parsed = ParsedDocument("docling", "2.115.0", (ParsedItem("image", "image", raw_text="OCR text", asset_bytes=b"image-bytes", asset_name="page.png", page_number=1),), raw_artifacts=(("docling_page_0001.json", b"{}"),))
            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=parsed), patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.write", return_value=1):
                ocr_only = service.ingest([str(source)])
            checkpoint_database = root / "indexes" / ocr_only["index_version"] / "checkpoints.sqlite3"
            connection = sqlite3.connect(checkpoint_database)
            try:
                row = connection.execute("SELECT checkpoint_json FROM local_parse_checkpoints").fetchone()
            finally:
                connection.close()
            self.assertIsNotNone(row)
            checkpoint = json.loads(row[0])
            source_asset = root / "assets" / checkpoint["items"][0]["asset_uri"]
            source_asset.unlink()
            with patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=parsed) as parse_page, patch("Agent.knowledge_base.multimodal.pipeline.StagedIndex.write", return_value=1):
                service.ingest([str(source)], allow_remote_data=True, max_images=1, reuse_local_from_index_version=ocr_only["index_version"])
            parse_page.assert_called_once()

    def test_local_checkpoint_requires_schema_and_matching_parser_contract(self) -> None:
        """不带 schema 或 Docling 版本不匹配的 checkpoint 不得被复用。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            store = AssetStore(root / "assets")
            entry = {"relative_path": "page.pdf", "content_hash": "a" * 64}
            parsed = ParsedDocument("docling", service._docling_version(), (ParsedItem("text", "paragraph", raw_text="text", page_number=1),))
            checkpoint = service._build_local_parse_checkpoint(root, "doc_" + "a" * 64, 1, parsed, entry, 1, store)
            checkpoint.pop("schema_version")
            service._write_local_parse_checkpoint(root, "doc_" + "a" * 64, 1, checkpoint)
            self.assertIsNone(service._reusable_local_parse_checkpoint(root, "doc_" + "a" * 64, 1, entry, 1, store))
            checkpoint["schema_version"] = "local-parse-v2"
            checkpoint["contract"]["parser"]["version"] = "mismatch"
            service._write_local_parse_checkpoint(root, "doc_" + "a" * 64, 1, checkpoint)
            self.assertIsNone(service._reusable_local_parse_checkpoint(root, "doc_" + "a" * 64, 1, entry, 1, store))

    def test_split_text_units_prefers_headings_and_keeps_target_size(self) -> None:
        """TXT/Markdown 切分必须优先标题/段落边界，目标约 800、最大 1200、重叠约 120。"""
        from Agent.knowledge_base.multimodal.parsers import TEXT_SPLIT, split_text_units
        self.assertEqual(TEXT_SPLIT["target_chars"], 800)
        self.assertEqual(TEXT_SPLIT["max_chars"], 1200)
        self.assertEqual(TEXT_SPLIT["overlap_chars"], 120)
        short = "短文不切分"
        self.assertEqual(split_text_units(short), [short])
        paragraphs = "\n\n".join(f"段落{index}：" + ("因果推理。" * 90) for index in range(1, 6))
        chunks = split_text_units(paragraphs)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk) <= TEXT_SPLIT["max_chars"] for chunk in chunks))
        self.assertTrue(any(len(chunk) >= TEXT_SPLIT["target_chars"] * 0.6 for chunk in chunks))
        # 每节约 900 字，超过 target；标题边界应阻止并入上一节
        heading_text = "# 标题一\n" + ("正文甲。" * 225) + "\n\n## 标题二\n" + ("正文乙。" * 225)
        headed = split_text_units(heading_text)
        self.assertGreaterEqual(len(headed), 2)
        self.assertTrue(any(chunk.lstrip().startswith("# 标题一") for chunk in headed))
        self.assertTrue(any(chunk.lstrip().startswith("## 标题二") or "标题二" in chunk[:40] for chunk in headed))
        hard = "字" * 2500
        hard_chunks = split_text_units(hard)
        self.assertGreaterEqual(len(hard_chunks), 2)
        self.assertTrue(all(len(chunk) <= TEXT_SPLIT["max_chars"] for chunk in hard_chunks))
        if len(hard_chunks) >= 2:
            overlap = min(TEXT_SPLIT["overlap_chars"], len(hard_chunks[0]), len(hard_chunks[1]))
            self.assertTrue(hard_chunks[0][-overlap:] == hard_chunks[1][:overlap] or hard_chunks[0][-40:] in hard_chunks[1])

    def test_parse_text_emits_multiple_paragraph_units(self) -> None:
        """长 TXT/Markdown 文件必须被切成多个 text/paragraph 单元，而不是单个超长单元。"""
        from Agent.knowledge_base.multimodal.parsers import TEXT_SPLIT
        with tempfile.TemporaryDirectory() as directory:
            body = "\n\n".join(f"## 段{index}\n" + ("因果。" * 100) for index in range(1, 5))
            source = Path(directory) / "notes.md"
            source.write_text(body, encoding="utf-8")
            parsed = parse_document(source, "docling")
            self.assertEqual(parsed.parser_name, "text")
            self.assertGreater(len(parsed.items), 1)
            self.assertTrue(all(item.modality == "text" and item.content_kind == "paragraph" for item in parsed.items))
            self.assertTrue(all(len(item.raw_text) <= TEXT_SPLIT["max_chars"] for item in parsed.items))
            joined = "\n".join(item.raw_text for item in parsed.items)
            self.assertIn("因果", joined)

    def test_manifest_includes_text_split_parameters(self) -> None:
        """文本切分参数必须写入 manifest.build_configuration，策略变化可生成新版本。"""
        from Agent.knowledge_base.multimodal.parsers import TEXT_SPLIT
        service = MultimodalKnowledgeBaseMaintenance()
        manifest = service._manifest([], allow_remote_data=False)
        self.assertEqual(manifest["build_configuration"]["text_split"], TEXT_SPLIT)

    def test_source_signature_mismatch_is_blocking(self) -> None:
        """伪装成 PDF 的图片必须在解析和外发前被拒绝。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "misleading.pdf"
            source.write_bytes(_png_bytes())
            issue = inspect_source(source)
            self.assertIsNotNone(issue)
            self.assertEqual(issue.code, "source_format_mismatch")
            self.assertTrue(issue.blocking)

    def test_page_routes_are_deterministic_and_mutually_exclusive(self) -> None:
        """空白、对象图片和整页回退必须产生唯一且稳定的本地路由。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "page.png"
            source.write_bytes(_png_bytes())
            local = ParsedDocument("builtin", "1", (ParsedItem("text", "paragraph", raw_text="正文", page_number=1),))
            picture = ParsedDocument("builtin", "1", (ParsedItem("image", "image", asset_bytes=_png_bytes(), asset_name="picture.png", page_number=1),))
            fallback = ParsedDocument("docling", "2", (ParsedItem("page", "page_render", asset_bytes=_png_bytes(), asset_name="page.png", page_number=1),))
            blank = ParsedDocument("builtin", "1")
            self.assertEqual(decide_page_route(source, 1, local).route, "local_objects")
            self.assertEqual(decide_page_route(source, 1, picture).route, "remote_pictures")
            self.assertEqual(decide_page_route(source, 1, fallback).route, "remote_page_fallback")
            self.assertEqual(decide_page_route(source, 1, blank).route, "blank_page")
            self.assertEqual(MultimodalKnowledgeBaseMaintenance._routed_page_items((picture.items[0], fallback.items[0]), "remote_page_fallback"), (fallback.items[0],))

    def test_page_preparation_keeps_unpaired_short_text(self) -> None:
        """质量门后的单页短正文不得因标题合并策略而静默丢失。"""
        service = MultimodalKnowledgeBaseMaintenance()
        items, filtered = service._prepare_page_items((ParsedItem("text", "paragraph", raw_text="短正文", page_number=1),))
        self.assertEqual(filtered, 1)
        self.assertEqual([item.raw_text for item in items], ["短正文"])

    def test_prepare_outbound_manifest_records_page_fallback_without_picture_duplicate(self) -> None:
        """整页 fallback 只能形成一个 page_render 外发记录，且携带质量门证据。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "page.png"
            source.write_bytes(_png_bytes())
            service = MultimodalKnowledgeBaseMaintenance(asset_root=root / "assets", index_root=root / "indexes", active_config=root / "active.json")
            parsed = ParsedDocument("docling", "2", (ParsedItem("page", "page_render", asset_bytes=_png_bytes(), asset_name="page.png", page_number=1),))
            output = root / "approved.json"
            with patch.object(service, "_remote_resource_allowed", return_value=True), patch("Agent.knowledge_base.multimodal.pipeline.parse_document_page", return_value=parsed):
                result = service.prepare_outbound_manifest([str(source)], output)
            records = service._read_outbound_manifest(output)
            self.assertEqual(result["outbound_image_count"], 1)
            self.assertEqual(records[0].route, "remote_page_fallback")
            self.assertEqual(records[0].quality_gate_version, "page-quality-v1")
            self.assertTrue(records[0].quality_summary)

if __name__ == "__main__":
    unittest.main()
